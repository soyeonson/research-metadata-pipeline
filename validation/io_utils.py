from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from .normalize import build_authority_map


def load_rules(path: str | Path) -> dict[str, Any]:
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def load_extracted_json(path: str | Path) -> dict[str, Any]:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def load_master_workbook(path: str | Path) -> tuple[list[dict[str, Any]], dict[str, str]]:
    wb = load_workbook(path, data_only=False)
    try:
        ws = wb["master_metadata"]
        headers = [cell.value for cell in ws[1]]
        records: list[dict[str, Any]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in row):
                continue
            records.append(dict(zip(headers, row)))

        authority_rows: list[tuple[object, object]] = []
        if "org_authority" in wb.sheetnames:
            aws = wb["org_authority"]
            for row in aws.iter_rows(min_row=2, values_only=True):
                if len(row) >= 2 and row[0] and row[1]:
                    authority_rows.append((row[0], row[1]))

        return records, build_authority_map(authority_rows)
    finally:
        wb.close()
