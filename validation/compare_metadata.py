from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .io_utils import load_extracted_json, load_master_workbook, load_rules
from .normalize import (
    clean_text,
    normalize_affiliation,
    normalize_oa,
    normalize_project_number,
    normalize_simple,
    split_semicolon,
)

STATUS_MATCH = "MATCH"
STATUS_MISMATCH = "MISMATCH"
STATUS_MISSING = "MISSING_IN_MASTER"
STATUS_REVIEW = "REVIEW"
STATUS_UNDETERMINED = "UNDETERMINED"


def _result(paper_id: str, field: str, master: Any, extracted: Any, status: str, note: str = "") -> dict[str, Any]:
    return {
        "paper_id": paper_id,
        "field": field,
        "master_value": "" if master is None else str(master),
        "extracted_value": "" if extracted is None else str(extracted),
        "status": status,
        "note": note,
    }


def compare_record(master: dict[str, Any], extracted: dict[str, Any], rules: dict[str, Any], authority_map: dict[str, str]) -> list[dict[str, Any]]:
    paper_id = clean_text(master.get("paper_id"))
    rows: list[dict[str, Any]] = []

    # Title / journal: deterministic normalized string comparison.
    for field in ("title", "journal"):
        m = master.get(field)
        e = extracted.get(field)
        if not clean_text(e):
            status = STATUS_UNDETERMINED
            note = "Extracted value unavailable"
        elif normalize_simple(m) == normalize_simple(e):
            status = STATUS_MATCH
            note = ""
        else:
            status = STATUS_MISMATCH
            note = "Normalized strings differ"
        rows.append(_result(paper_id, field, m, e, status, note))

    # OA: master uses Y or blank in this synthetic example.
    m_oa = normalize_oa(master.get("oa"))
    e_oa = normalize_oa(extracted.get("open_access"))
    if e_oa is None:
        rows.append(_result(paper_id, "oa", m_oa, e_oa, STATUS_UNDETERMINED, "Source did not establish OA"))
    elif m_oa is None and e_oa == "Y":
        rows.append(_result(paper_id, "oa", m_oa, e_oa, STATUS_MISSING, "Extracted source indicates OA but master is blank"))
    elif m_oa == e_oa:
        rows.append(_result(paper_id, "oa", m_oa, e_oa, STATUS_MATCH))
    elif m_oa == "Y" and e_oa == "N":
        rows.append(_result(paper_id, "oa", m_oa, e_oa, STATUS_REVIEW, "Conflicting OA signal requires review"))
    else:
        rows.append(_result(paper_id, "oa", m_oa, e_oa, STATUS_MISMATCH))

    # Project numbers: normalize dashes/prefixes then compare as sets.
    project_rules = rules.get("project_number", {})
    prefixes = project_rules.get("strip_prefixes", [])
    m_projects = {normalize_project_number(x, prefixes) for x in split_semicolon(master.get("project_numbers"))}
    e_projects = {normalize_project_number(x, prefixes) for x in extracted.get("project_numbers", []) if clean_text(x)}
    m_projects.discard("")
    e_projects.discard("")
    if not e_projects:
        rows.append(_result(paper_id, "project_numbers", sorted(m_projects), sorted(e_projects), STATUS_UNDETERMINED, "No project number extracted"))
    elif m_projects == e_projects:
        rows.append(_result(paper_id, "project_numbers", sorted(m_projects), sorted(e_projects), STATUS_MATCH))
    else:
        rows.append(_result(paper_id, "project_numbers", sorted(m_projects), sorted(e_projects), STATUS_MISMATCH, "Normalized project-number sets differ"))

    # Affiliation: article-level unique set for demo purposes.
    m_aff = {normalize_affiliation(x, authority_map) for x in split_semicolon(master.get("affiliations"))}
    e_aff_raw: list[str] = []
    for author in extracted.get("authors", []):
        e_aff_raw.extend(author.get("affiliations", []))
    e_aff = {normalize_affiliation(x, authority_map) for x in e_aff_raw if clean_text(x)}
    m_aff.discard("")
    e_aff.discard("")
    if not e_aff:
        rows.append(_result(paper_id, "affiliations", sorted(m_aff), sorted(e_aff), STATUS_UNDETERMINED, "No affiliation extracted"))
    elif {normalize_simple(x) for x in m_aff} == {normalize_simple(x) for x in e_aff}:
        rows.append(_result(paper_id, "affiliations", sorted(m_aff), sorted(e_aff), STATUS_MATCH))
    else:
        rows.append(_result(paper_id, "affiliations", sorted(m_aff), sorted(e_aff), STATUS_MISMATCH, "Affiliation sets differ after authority normalization"))

    # Author count: simple quality signal, not identity matching.
    m_count = master.get("author_count")
    e_count = len(extracted.get("authors", []))
    if m_count in (None, ""):
        rows.append(_result(paper_id, "author_count", m_count, e_count, STATUS_UNDETERMINED, "Master author count unavailable"))
    elif int(m_count) == e_count:
        rows.append(_result(paper_id, "author_count", m_count, e_count, STATUS_MATCH))
    else:
        rows.append(_result(paper_id, "author_count", m_count, e_count, STATUS_MISMATCH, "Author counts differ; manual author-level review recommended"))

    return rows


def compare_all(master_records: list[dict[str, Any]], extracted_payload: dict[str, Any], rules: dict[str, Any], authority_map: dict[str, str]) -> list[dict[str, Any]]:
    extracted_index = {clean_text(p.get("paper_id")): p for p in extracted_payload.get("papers", [])}
    results: list[dict[str, Any]] = []
    for master in master_records:
        paper_id = clean_text(master.get("paper_id"))
        extracted = extracted_index.get(paper_id)
        if extracted is None:
            results.append(_result(paper_id, "record", "present", "missing", STATUS_UNDETERMINED, "No extracted record matched paper_id"))
            continue
        results.extend(compare_record(master, extracted, rules, authority_map))
    return results


def write_report(results: list[dict[str, Any]], output_path: str | Path) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "validation_results"
    headers = ["paper_id", "field", "master_value", "extracted_value", "status", "note"]
    ws.append(headers)

    for row in results:
        ws.append([row[h] for h in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    status_fills = {
        STATUS_MATCH: "E2F0D9",
        STATUS_MISMATCH: "F4CCCC",
        STATUS_MISSING: "FCE5CD",
        STATUS_REVIEW: "FFF2CC",
        STATUS_UNDETERMINED: "D9D2E9",
    }
    for row in ws.iter_rows(min_row=2):
        status_cell = row[4]
        status_cell.fill = PatternFill("solid", fgColor=status_fills.get(str(status_cell.value), "FFFFFF"))
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {"A": 14, "B": 20, "C": 44, "D": 44, "E": 22, "F": 52}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)
    wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare master metadata with structured extracted metadata.")
    parser.add_argument("--master", required=True)
    parser.add_argument("--extracted", required=True)
    parser.add_argument("--rules", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    master_records, authority_map = load_master_workbook(args.master)
    extracted = load_extracted_json(args.extracted)
    rules = load_rules(args.rules)
    results = compare_all(master_records, extracted, rules, authority_map)
    write_report(results, args.output)

    counts = Counter(row["status"] for row in results)
    print(f"Wrote {len(results)} validation rows to {args.output}")
    print(dict(counts))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
